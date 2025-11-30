import json
import os
from docx import Document

def docx_to_json(docx_path, json_path):
    if not os.path.exists(docx_path):
        print(f"❌ File not found: {docx_path}")
        return

    document = Document(docx_path)
    data = []
    
    # Strategy: Try to identify Q&A pairs
    # We'll look for paragraphs. If we find a pattern like "Question:" or "Q:", we treat it as a question.
    
    current_item = {}
    
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
            
        # Check for Question
        if text.lower().startswith(('question:', 'q:', 'question', 'q ')):
            # Save previous item if exists
            if 'question' in current_item and 'answer' in current_item:
                data.append(current_item)
                current_item = {}
            
            # Clean up text
            clean_text = text.split(':', 1)[-1].strip() if ':' in text else text
            current_item['question'] = clean_text
            
        # Check for Answer
        elif text.lower().startswith(('answer:', 'a:', 'answer', 'a ')):
            clean_text = text.split(':', 1)[-1].strip() if ':' in text else text
            current_item['answer'] = clean_text
            
        # If we have a question but no answer yet, append to question (multiline)
        elif 'question' in current_item and 'answer' not in current_item:
            current_item['question'] += " " + text
            
        # If we have an answer, append to answer (multiline)
        elif 'answer' in current_item:
            current_item['answer'] += " " + text

    # Add the last item
    if 'question' in current_item and 'answer' in current_item:
        data.append(current_item)

    # Fallback: If no Q&A structure found, just dump paragraphs
    if not data:
        print("⚠️ No structured Q&A found. Converting paragraphs to simple list.")
        data = [{"text": p.text.strip()} for p in document.paragraphs if p.text.strip()]

    # Save to JSON
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=4, ensure_ascii=False)
        
    print(f"✅ Converted {len(data)} items to {json_path}")

if __name__ == "__main__":
    docx_to_json('data/data.docx', 'data/data.json')
